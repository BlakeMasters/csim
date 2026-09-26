"""Read-only format and integrity smoke for local XLSX and MATLAB v5 inputs.

This reports structure and missingness. It does not map biological observations,
choose splits, filter rows, or establish model validity.
"""

from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import importlib.metadata
import json
from pathlib import Path
import sys
import zipfile

import numpy as np
import openpyxl
from scipy.io import loadmat, whosmat


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def inspect_xlsx(path: Path) -> dict:
    with zipfile.ZipFile(path) as archive:
        bad_member = archive.testzip()
        if bad_member is not None:
            raise ValueError(f"ZIP member failed CRC: {bad_member}")

    book = openpyxl.load_workbook(
        path, read_only=True, data_only=False, keep_links=False
    )
    try:
        sheets = []
        for sheet in book:
            rows = sheet.iter_rows()
            first = next(rows, ())
            headers = [cell.value for cell in first]
            repeated = {
                str(name): [index + 1 for index, value in enumerate(headers) if value == name]
                for name, count in Counter(headers).items()
                if name is not None and count > 1
            }
            observations = 0
            formula_cells = sum(cell.data_type == "f" for cell in first)
            error_cells = sum(cell.data_type == "e" for cell in first)
            for row in rows:
                if any(cell.value is not None and str(cell.value).strip() for cell in row):
                    observations += 1
                formula_cells += sum(cell.data_type == "f" for cell in row)
                error_cells += sum(cell.data_type == "e" for cell in row)
            sheets.append(
                {
                    "name": sheet.title,
                    "columns": sheet.max_column,
                    "nonempty_rows_after_header": observations,
                    "duplicate_headers_1based": repeated,
                    "formula_cells": formula_cells,
                    "error_cells": error_cells,
                }
            )
    finally:
        book.close()
    return {"name": path.name, "bytes": path.stat().st_size, "sha256": sha256(path), "sheets": sheets}


def inspect_mat(path: Path) -> dict:
    with path.open("rb") as stream:
        header = stream.read(128)
    if not header.startswith(b"MATLAB 5.0 MAT-file"):
        raise ValueError("Expected a MATLAB v5/v7 binary MAT file")

    descriptions = whosmat(str(path))
    if not descriptions:
        raise ValueError("MAT file has no top-level variables")
    arrays = loadmat(
        str(path),
        variable_names=[name for name, _, _ in descriptions],
        verify_compressed_data_integrity=True,
    )
    variables = []
    for name, shape, matlab_class in descriptions:
        array = arrays[name]
        numeric = np.issubdtype(array.dtype, np.number)
        missing = None
        if numeric:
            missing = {
                "nan": int(np.isnan(array).sum()),
                "positive_infinity": int(np.isposinf(array).sum()),
                "negative_infinity": int(np.isneginf(array).sum()),
            }
        variables.append(
            {
                "name": name,
                "shape": list(shape),
                "matlab_class": matlab_class,
                "dtype": str(array.dtype),
                "nonfinite": missing,
            }
        )
    return {"name": path.name, "bytes": path.stat().st_size, "sha256": sha256(path), "variables": variables}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path("data"))
    args = parser.parse_args()
    root = args.root.resolve()
    if not root.is_dir():
        parser.error(f"Data directory does not exist: {root}")

    files = sorted(path for path in root.iterdir() if path.is_file())
    report = {
        "schema": "cellsim_v2.local-data-reader-check.v1",
        "root": str(root),
        "python": sys.version.split()[0],
        "packages": {
            name: importlib.metadata.version(name)
            for name in ("numpy", "scipy", "openpyxl", "ocura-oss")
        },
        "files": [],
        "problems": [],
    }
    for path in files:
        try:
            if path.suffix.lower() == ".xlsx":
                report["files"].append(inspect_xlsx(path))
            elif path.suffix.lower() == ".mat":
                report["files"].append(inspect_mat(path))
            else:
                report["problems"].append({"file": path.name, "error": "Unsupported extension"})
        except (OSError, ValueError, TypeError, zipfile.BadZipFile) as exc:
            report["problems"].append({"file": path.name, "error": f"{type(exc).__name__}: {exc}"})
    report["status"] = "passed" if not report["problems"] and files else "failed"
    print(json.dumps(report, sort_keys=True, allow_nan=False))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
